"""Build the report from the supplied template while preserving its package furniture.
Run with the bundled document Python. Pass --template if the reference lives elsewhere.
"""
from pathlib import Path
from copy import deepcopy
import argparse, json, zipfile, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser();parser.add_argument('--template',default='/Users/aruispan/Downloads/Report template (2).docx');parser.add_argument('--submission-date',default='');args=parser.parse_args()
reference=Path(args.template);d=Document(reference)
metrics=json.loads((root/'docs/evidence/metrics.json').read_text())
refs=json.loads((root/'docs/references.json').read_text())
# Keep cover paragraphs and the source section properties, remove instructional body slots.
body=d._element.body
for child in list(body)[13:]:
 if child.tag!=qn('w:sectPr'):body.remove(child)
for p in d.paragraphs:
 if '[Insert Project Title]' in p.text:
  p.runs[-1].text=' Baikonur Rocket Launch'
  try:p.style=d.styles['Title']
  except KeyError:pass
  for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
 if '[Insert Names and group]' in p.text:p.runs[-1].text=' Akbope Bakytkeldy, SE2426\nMoldir Omar, SE2426\nAyaulym Talgat, SE2427\nTalant Altaikhan, SE2425'
 if '[Insert Date]' in p.text:p.runs[-1].text=' '+args.submission_date

pages=[]
def p(text='',bold=False,center=False,keep=False,size=12):
 q=d.add_paragraph();q.paragraph_format.space_after=Pt(8);q.paragraph_format.line_spacing=278/240
 q.paragraph_format.keep_with_next=keep
 if center:q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=q.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
 return q

def new(title):
 d.add_page_break();pages.append(title);p(title,True,keep=True)
def sub(title):p(title,True,keep=True)
def eq(text):
 import re
 q=p(center=True);m=OxmlElement('m:oMath')
 def run(value):
  r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r
 for token in re.split(r'([A-Za-zΣ]+_[A-Za-z]+|\^[A-Za-z0-9]+)',text):
  if '_' in token:
   base,sub=token.split('_');n=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(run(base));b=OxmlElement('m:sub');b.append(run(sub));n.extend([e,b]);m.append(n)
  elif token.startswith('^'):
   n=OxmlElement('m:sSup');e=OxmlElement('m:e');e.append(m[-1]);b=OxmlElement('m:sup');b.append(run(token[1:]));n.extend([e,b]);m.append(n)
  elif token:m.append(run(token))
 q._p.append(m)
def fig(file,caption,width=6.35):
 q=p(center=True,keep=True);q.add_run().add_picture(str(root/'docs/evidence'/file),width=Inches(width))
 p(caption,center=True)
def table(headers,rows,widths=None):
 t=d.add_table(rows=1, cols=len(headers));t.autofit=False
 total=6.5
 if widths is None:widths=[total/len(headers)]*len(headers)
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for cell,text in zip(t.rows[0].cells,headers):cell.text=text
 for row in rows:
  cells=t.add_row().cells
  for cell,text in zip(cells,row):cell.text=str(text)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for i,cell in enumerate(row.cells):
   cell.width=Inches(widths[i])
   for q in cell.paragraphs:
    q.paragraph_format.space_after=Pt(5);q.paragraph_format.space_before=Pt(3)
    for r in q.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=ri==0
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'888888');borders.append(e)
 t._tbl.tblPr.append(borders)
 p()
 return t

new('Table of Contents')
for line in [
'1 Introduction ........................................................................ 3',
'2 Literature Review .............................................................. 4',
'   2.1 Particle and fire representations .................................. 4',
'   2.2 Turbulence and rendering cost ..................................... 5',
'   2.3 Materials and atmospheric appearance ........................ 6',
'3 Project Description ............................................................ 7',
'4 Technical Implementation .................................................. 8',
'   4.1 Tools and technologies used ....................................... 8',
'   4.2 Development process ................................................ 8',
'   4.3 Algorithms and mathematical model ........................... 9',
'5 Graphics and Rendering ................................................... 10',
'   5.1 Scene construction and lighting ................................ 10',
'   5.2 Shading and texture sampling ................................... 11',
'   5.3 Particle animation and separation .............................. 12',
'6 Testing and Evaluation ..................................................... 13',
'7 Conclusion and Future Work ............................................. 14',
'8 References ....................................................................... 15',
'9 AI Use Note ..................................................................... 18',
'Appendix A Course connections and defense ..................... 19']:
 p(line)
p('Deliverables',True)
p('The accompanying source package contains the C++ application, GLSL shaders, bundled dependencies, tests, build instructions and evidence logs. The presentation contains ten slides. The captioned demonstration runs for 3 minutes 20 seconds at 1920 by 1080 pixels.')

new('1 Introduction')
p('Baikonur Rocket Launch is an interactive computer graphics application that presents a countdown, engine ignition, ascent and booster separation from a stylized launch complex. The project combines the course pipeline into one native C++ application: indexed geometry, programmable shaders, matrix transformations, lighting, texture sampling and interactive viewing.')
p('A rocket launch makes the relationship between these techniques visible. The gantry and ground provide stable spatial references. The rocket body and four boosters demonstrate hierarchical transformations. Ignition introduces a moving point light and two different particle materials. A chase camera connects the flight path to the viewing transformation, while an altitude-dependent sky shows how a simulation parameter can drive a fragment shader.')
p('The application is intended for classroom demonstration and explanation. The rocket and site are original procedural approximations inspired by the selected Baikonur topic. Their dimensions, trajectory and timing are compressed to make the complete sequence easy to inspect. The result is a graphics study rather than an aerospace simulator or an architectural reconstruction of a specific launch pad.')
sub('1.1 Expected outcome')
p('The principal outcome is a repeatable, interactive scene that exposes its graphics features through keyboard, mouse and on-screen controls. Users can inspect the launch site, switch lighting, view wireframe geometry, launch the rocket, follow its ascent, separate boosters, pause the mission and reset it. A visible FPS counter and exported frame-time logs support performance evaluation.')
p('The implementation includes all three topic-specific stretch features: booster detachment with independent motion, ignition camera shake, and stars that become visible at altitude. The report separates these implemented visual effects from advanced research techniques discussed as alternatives. It also records which validation tasks remain specific to the final teaching-assistant environment.')
sub('1.2 Relation to the course')
p('The implementation follows the concepts in the supplied lectures on transformations, viewing, shading, texturing, hierarchical models and procedural methods [16]. The final assessment requires OpenGL 3.3 or newer in core profile. Consequently, the project uses GLFW and a GLAD loader with explicit matrices and GLSL rather than the historical fixed-function examples that appear in some lecture slides.')

new('2 Literature Review')
sub('2.1 Particle and fire representations')
p('The review compares fifteen journal articles [1]-[15] from ACM Transactions on Graphics and Computer Graphics Forum. It focuses on decisions relevant to a small interactive launch scene: visual control, simulation cost, material representation and atmospheric appearance. The cited papers establish alternatives and limitations; the project does not claim to implement every reviewed method.')
p('Reeves [1] represents fuzzy phenomena through many particles whose birth, motion, appearance and death vary over time. This representation is a direct conceptual fit for exhaust and smoke because a cloud has no single stable polygonal boundary. The project adopts a finite particle lifecycle with randomized initial velocity, size and lifetime. Its camera-facing quads differ from a full participating-volume renderer but keep the lifecycle visible and inexpensive.')
p('Nguyen, Fedkiw and Jensen [5] model fuel and hot gaseous products using fluid dynamics and render the resulting fire through a physically based treatment. This connects flame motion, buoyancy and radiance more closely than independent particles can. Such a system would improve combustion realism, but its field representation and numerical solution would substantially expand the scope of a course project. The implemented flame therefore uses age-based color, downward velocity and additive blending.')
p('Lamorlette and Foster [6] emphasize artistic control of flame behavior through structural and stochastic mechanisms. Their approach shows that controllable motion is itself an important production objective. In this project, explicit emission rate, lifetime, speed and color parameters serve a similar design objective at a much simpler level. No spline-based flame skeleton or production flame system from that paper is implemented.')
p('Rasmussen and colleagues [7] combine fluid-generated motion with additional turbulent detail to produce large-scale smoke. Their work explains why a convincing plume needs structure at several spatial scales. The project instead uses expanding billboards with outward drift near the pad. This is sufficient to distinguish smoke from fire, but it cannot reproduce coherent rolling vortices or light attenuation through a dense volume.')
sub('Design implication')
p('These papers support a deliberate scope decision: maintain a clear and bounded particle lifecycle first, then identify fluid simulation and richer turbulent structure as future work. The visible effect must remain responsive to launch, pause and reset commands. This favors a compact CPU pool and GPU instancing over a large volumetric simulation.')

new('2.2 Turbulence and rendering cost')
p('Losasso, Gibou and Fedkiw [8] use an octree to concentrate simulation resolution where detail matters. This reduces the waste of uniformly fine grids but introduces adaptive data structures and pressure-solver complexity. The launch project has a different bottleneck and scale: a small, bounded particle set. It therefore shares geometry and batches particle rendering rather than allocating a fluid grid.')
p('Selle, Rasmussen and Fedkiw [9] combine vortex particles with an Eulerian representation to recover turbulent detail that can be lost in grid-based advection. Their method treats particles as carriers of rotational flow information, whereas the project particles are independent visual elements. The distinction is important: the implemented smoke has plausible drift and expansion, but it does not preserve vorticity or satisfy incompressible flow equations.')
p('Bridson, Hourihan and Nordenstam [10] derive procedural velocity fields using curl noise, enabling incompressible turbulent motion and boundary-aware behavior without a full simulation. This offers a closer extension path for the current application than an adaptive fluid solver. A future version could replace the simple outward smoke velocity with a curl-noise field. The current release does not label its sinusoidal ground pattern or random particle velocities as curl noise.')
p('Kim and colleagues [11] separate large-scale motion from high-frequency turbulence using wavelets. This separation helps artists add detail without recomputing every aspect of the base flow. It also suggests a useful architecture principle: keep mission progression independent from appearance detail. The project follows that separation through distinct Mission and Particles modules, although it does not implement wavelet synthesis.')
p('Cook [3] studies sampling and aliasing in computer graphics and shows how sampling patterns influence visible artifacts. The relevant lesson here is that increased geometric detail alone cannot prevent shimmer. The application requests four-sample MSAA for ordinary window rendering and uses mipmapped minification for its concrete texture. These are distinct practical filters, not an implementation of Cook\'s stochastic ray-tracing method. Video capture also renders with four-sample MSAA, then resolves to a texture before readback.')
sub('Comparative assessment')
p('Adaptive grids [8], hybrid vorticity [9] and wavelet detail [11] offer more physical or multiscale structure than independent billboards. Curl noise [10] offers procedural control with lower integration cost. The current implementation prioritizes understandable data flow and stable interaction, using at most 3,600 particles and two instanced particle draws. The research methods remain specific, testable directions for future improvement.')

new('2.3 Materials and atmospheric appearance')
p('Cook and Torrance [2] describe a reflectance model that accounts for material-dependent angular behavior. Compared with that microfacet approach, Blinn-Phong offers a smaller set of terms that directly matches the course assessment: ambient, diffuse and specular. The project uses the latter and assigns different shininess values to metal-like rocket parts and matte structures. It does not claim physically based material calibration or energy conservation.')
p('Perlin [4] improves the construction of procedural noise, while Lagae and colleagues [15] compare noise families, filtering and their uses in graphics. These works explain why continuity and frequency content matter when procedural patterns are sampled. The implementation uses an original low-frequency trigonometric steppe pattern and a deterministic random concrete image. Neither is described as Perlin noise. More sophisticated noise would improve the natural appearance of the terrain and smoke.')
p('Bruneton and Neyret [12] precompute atmospheric scattering so a renderer can evaluate views from the ground to space efficiently. This approach relates appearance to physical scattering and viewing geometry. The application uses a much simpler blend between a warm horizon and a dark sky as the rocket gains altitude. It meets the visual transition required by the topic, but it cannot reproduce physically correct aerial perspective, planetary shadows or multiple scattering.')
p('Hosek and Wilkie [13] fit an analytic spectral sky model that improves important daylight conditions, including sunsets. Its compact evaluation makes it relevant to real-time graphics, but it targets a more principled radiance distribution than the project\'s artistic gradient. Hillaire [14] addresses scalable atmosphere rendering with compact lookup-table parameterizations and a multiple-scattering approximation. That method is a stronger future option for continuous ground-to-space views.')
sub('Synthesis and source quality')
p('The final design uses simple course-level shading and particles while retaining an architecture that can accept richer material, noise and atmosphere models. The twelve ACM Transactions on Graphics articles and three Computer Graphics Forum articles are journal publications, including journal issues associated with graphics conferences. Their DOI metadata was checked against publisher records. Both venues are Q1 in the 2024 SCImago Computer Graphics and Computer-Aided Design category, using Scopus data as of March 2025 [19]. This is a venue-level classification; the submission checklist records the evidence and ranking convention.')

new('3 Project Description')
sub('3.1 Objective')
p('Build a native C++17 application that makes the full graphics pipeline observable through a polished, interactive Baikonur-themed launch sequence. The target is at least 30 FPS at 1280 by 720 on the teaching-assistant machine, with a preference for 60 FPS or more. Results measured on the development machine are reported in Section 6.')
sub('3.2 Problem statement')
p('The central problem is coordinating multiple representations in one frame. Rocket parts require consistent parent transforms, smoke requires transparency ordering, lighting requires normals in the correct coordinate space, and the camera must follow an object whose position changes continuously. These systems must remain synchronized when the user pauses, changes playback speed or resets the mission.')
sub('3.3 Scope')
table(['Area','Implemented behavior'],[
('Scene','Steppe, pad, flame trench, gantry, rocket, four boosters, buildings, tanks, radome, roads and lights.'),
('Mission','Eight-second countdown, ignition at T-2, liftoff at T+0, automatic separation at T+24, manual separation after T+5.'),
('Graphics','Indexed meshes, original GLSL, inverse-transpose normals, three light contributions, mipmapped concrete and procedural sky.'),
('Interaction','Keyboard and mouse orbit, zoom, chase mode, pause, reset, speed control, light/night toggles, wireframe and screenshots.'),
('Stretch','Independent booster trajectories, ignition shake, altitude-dependent stars.')],[1.05,5.45])
p('The site and vehicle are stylized. Flight dynamics, combustion chemistry, Earth curvature, real telemetry, shadow mapping and collision with the gantry are outside scope. Booster motion uses a visual ballistic arc and a ground-height clamp, not a rigid-body collision solver. The sky transition occurs at compressed demonstration altitudes.')

new('4 Technical Implementation')
sub('4.1 Tools and technologies used')
p('The application uses C++17, CMake, GLFW 3.4, the GLAD loader included with GLFW, GLM 1.0.1 and Dear ImGui 1.91.8. Dependencies are included as source and build with the project. Scene shaders use GLSL 330 core. On the test Mac, the driver provides an OpenGL 4.1 core context while the application uses the required 3.3 feature set.')
table(['Module','Responsibility'],[
('main','Window/context, input, elapsed time, render order, capture and frame logs.'),
('Shader / Mesh / Texture','Compile and link GLSL, own VAO/VBO/EBO resources, upload geometry once, generate the concrete texture and mip chain.'),
('Mission / Camera','Mission state, altitude, booster transforms, orbit/chase viewing and ignition shake.'),
('Scene / Particles','Static scene and rocket hierarchy; particle lifetime updates and instanced billboard passes.'),
('UI / RenderTarget','Controls, telemetry and captions; fixed-resolution recording target displayed in the application window.')],[1.6,4.9])
sub('4.2 Development process')
p('Development proceeded through a buildable pipeline, primitive geometry, mission transforms, lighting and textures, particles, interaction, visual validation and documentation. Automated tests isolate the simulation from the graphics context. Render captures check behavior that numerical tests cannot establish, including transparency, camera composition and UI placement.')
p('The initial implementation and documents received substantial AI assistance, detailed in Section 9. The package does not claim that the four students independently wrote particular modules or maintained a multiweek commit history. Proposed review and speaking ownership appears in the appendix. Members must record their actual review, changes and test results in genuine commits before submission.')

new('4.3 Algorithms and mathematical model')
sub('Transformation pipeline and hierarchy')
eq('p_clip = P V M_parent M_local p_object')
p('The application uses column vectors. Local scaling shapes each primitive, local translation positions it, and the rocket parent applies ascent. The view matrix comes from the camera eye, target and up vector. A perspective projection uses a 48 degree vertical field of view, the framebuffer aspect ratio, a 0.1 near plane and a 20,000 far plane.')
eq('n_world = normalize(transpose(inverse(mat3(M))) n_object)')
p('The normal matrix preserves perpendicularity under nonuniform scaling. Interpolated normals are normalized again in the fragment shader. Mesh UVs and world positions pass from the vertex stage to fragment evaluation.')
sub('Mission and booster motion')
eq('h(t) = 0.55 max(t, 0)^2 ;   v(t) = 1.1 max(t, 0)')
p('The mission clock begins at -8 seconds. At separation time ts, each booster preserves its world position and initial upward velocity. For elapsed time u = t - ts, radial distance grows as 1.72 + 1.8u and vertical position becomes 3.8 + h(ts) + 1.1tsu - 1.8u². Rotation increases at 0.35u radians. Continuity at u = 0 is checked by an automated matrix comparison.')
sub('Particle lifecycle and rendering order')
p('Each particle stores position, velocity, age, lifetime, size and material type. Age increases by elapsed time and expired particles are removed. An emission accumulator generates 700 particles per simulation second with a persistent 2:1 fire-to-smoke ratio. The pool is capped at 3,600. Static quad data uploads once; a compact per-instance position/size and color buffer updates each frame.')
p('The frame order is sky, opaque geometry, sorted smoke, additive fire and UI. Smoke sorts back to front using camera distance and blends with source alpha. Fire uses additive blending. Both particle passes keep depth testing against opaque geometry and disable depth writes. A soft radial fragment mask removes the square billboard boundary.')

new('5 Graphics and Rendering')
sub('5.1 Scene construction and lighting')
fig('day.png','Fig. 1 Launch complex with procedural geometry and live telemetry',5.5)
p('Figure 1 shows the shared primitive meshes assembled into distinct scene objects. The rocket transform parents the body, nose, bands and boosters. The orange gantry bracing and turquoise vehicle panels establish a consistent palette.')
fig('night.png','Fig. 2 Night lighting with the toggled gantry point light',5.5)
p('Figure 2 shows the scene under a different lighting mode. The renderer combines a directional sun, a blue gantry point light and an orange engine point light. The engine light follows the rocket during ascent.')

new('5.2 Shading and texture sampling')
eq('I = k_a C + Σ_j A_j L_j [C max(N · L,0) + k_s max(N · H,0)^s]')
p('Here C is the surface color, H is the normalized sum of light and view directions, s is shininess, and A is point-light attenuation. Specular contribution is zero when the light lies behind the surface. The shader adds emission for selected lamps and applies a display gamma approximation after shading.')
fig('wireframe.png','Fig. 3 Wireframe mode exposes indexed triangles in the scene',5.0)
p('Figure 3 exposes triangle topology. Wireframe mode draws unique edges with GL_LINES. Both modes use VAOs, vertex buffers and element buffers, with no fixed-function drawing calls.')
fig('texture.png','Fig. 4 Repeating concrete UVs and procedural steppe detail',5.0)
p('Figure 4 shows the concrete pad with a deterministic 1024 by 1024 RGB texture. The sampler repeats UVs, uses linear magnification and trilinear mipmapped minification. The steppe uses a separate low-frequency shader pattern. The concrete grid is a stylized material and does not preserve a fixed real-world slab size across differently scaled objects.')

new('5.3 Particle animation and separation')
fig('ignition.png','Fig. 5 Ignition combines fire, smoke and a moving engine light',5.5)
p('Figure 5 shows warm additive flame and expanding translucent smoke. Ignition also applies a short, decaying camera offset. The K key disables that motion. Pause freezes the mission and particles while allowing independent camera movement.')
fig('separation.png','Fig. 6 Detached boosters against the altitude-dependent star field',5.5)
p('Figure 6 shows the topic\'s final visual event. The shader blends the warm sky toward darkness using clamp(altitude / 450, 0, 1). A direction-based procedural star field fades in at altitude. Booster detachment switches from parent-relative motion to independent world transforms. These are artistic visual effects at compressed altitudes.')

new('6 Testing and Evaluation')
sub('6.1 Testing methodology')
p('CTest runs a standalone simulation test executable. It checks grounded idle state, ignition timing, pause, ascent, separation continuity, outward booster motion, equivalent progression at 30 and 60 updates per second, reset, rejection of premature separation, automatic separation, bounded atmosphere values and finite booster transforms. The automated suite passes. Rendering tests compile every shader and check OpenGL errors after each frame. No independent student user study has been conducted.')
sub('6.2 Performance evaluation')
p('Hardware: Apple M1 with 8 GiB unified memory, macOS 26.3.1. Compiler: AppleClang 21.0.0. Release build, OpenGL 4.1 Metal driver 90.5. Measurements below use 1280 by 720, four-sample window MSAA, VSync disabled and no recording. Each 12-second run excludes its first two seconds. Mean FPS is sample count divided by the sum of measured frame times; p95 reports the 95th percentile frame time.')
rows=[]
for r in metrics['runs']:rows.append((r['scene'],f"{r['mean_fps']:.1f}",f"{r['p95_ms']:.2f}",str(r['max_particles'])))
table(['Scene','Mean FPS','p95 ms','Max particles'],rows,[2.4,1.3,1.2,1.6])
p('These are local frame-loop throughput measurements, including CPU work and buffer swaps. They are not GPU-only timings and do not establish performance on the TA machine. Occasional frame outliers remain; these averages do not guarantee every frame meets the target. Raw CSV data and exact commands accompany the project. The edited demonstration uses 3840 by 2160 renderer captures, downsampled to 1920 by 1080 with a visible FPS counter and no playback speed change.')
p('Static geometry and the concrete mip chain upload once. Particle rendering uses two instanced draws. The allocated instance buffer is 115,200 bytes (3,600 instances times 32 bytes). The particle pool is bounded; total process/GPU memory was not instrumented, so no unsupported total-memory figure is reported.')
sub('6.3 Issues and debugging')
p('Validation found a repeated GLAD implementation include, a high-FPS bias in the fire/smoke mix, overlapping help text, and an incompatible multisample framebuffer blit in video capture. Fixes isolated the loader definition, made the emission ordinal persistent, repositioned help, and added a compatible multisample resolve before capture/display. Scripted demo input is locked to avoid unrelated user input changing its timeline.')

new('7 Conclusion and Future Work')
p('The project integrates a complete programmable graphics pipeline into an interactive launch scene. It demonstrates hierarchical modeling, per-fragment lighting, UV texture sampling, particle blending, camera control and real-time telemetry. All three Baikonur stretch features are implemented and appear in the demonstration. Automated simulation tests and local rendering checks support the implementation, while the final TA-machine and team defense checks remain separate responsibilities.')
sub('7.1 Limitations')
p('The animation follows a visual trajectory rather than a force-based rocket model. The site uses repeated primitives, the terrain is a large flat surface, and the sky is an artistic gradient. Smoke does not collide with the gantry or cast volumetric shadows. The current blend ordering uses particle-center distance, which is an approximation for overlapping billboards. At extreme altitude or camera distance, the finite terrain and long projection range can expose visual limits.')
p('A long application stall is clamped to a 0.1-second simulation update, preventing a sudden visual jump at the cost of falling behind wall time during exceptional delays. Four-sample MSAA smooths edges in both the window and the recording. Retina rendering, dense curved meshes and a high-resolution TrueType font atlas improve sharpness. Performance varies with display synchronization, framebuffer size, capture and background load. The measurements here do not replace a fresh build and run on the actual assessment machine.')
sub('7.2 Prioritized future work')
p('The first graphics extension would be directional shadow mapping to anchor the rocket and gantry more convincingly. A second would replace independent smoke drift with a curl-noise velocity field [10] and depth-aware soft particles to improve intersections. A third would introduce a curved ground model and a physical atmosphere approximation based on [12]-[14]. These additions should be profiled individually before increasing visual complexity.')
p('The strongest software improvements would be a more general scene-node representation, finer resource-error cleanup, and platform CI for Windows and Linux in addition to the tested Mac build. A structured user evaluation could measure whether observers can identify the major launch events and operate the controls without assistance. Each improvement should retain the project\'s central teaching purpose: making the graphics pipeline explainable.')

new('8 References')
p('Journal articles use IEEE-style numbering. DOI records identify the journal versions; conference-associated journal issues are not counted twice.')
for r in refs[:6]:p(f"[{r['id']}] {r['authors']}, \"{r['title']},\" {r['venue']}, vol. {r['volume']}, no. {r['issue']}, pp. {r['pages']}, {r['year']}. https://doi.org/{r['doi']}")
new('8 References continued')
for r in refs[6:12]:p(f"[{r['id']}] {r['authors']}, \"{r['title']},\" {r['venue']}, vol. {r['volume']}, no. {r['issue']}, {r['pages'] if 'Article' in r['pages'] else 'pp. '+r['pages']}, {r['year']}. https://doi.org/{r['doi']}")
new('8 References continued')
for r in refs[12:]:p(f"[{r['id']}] {r['authors']}, \"{r['title']},\" {r['venue']}, vol. {r['volume']}, no. {r['issue']}, {r['pages'] if 'Article' in r['pages'] else 'pp. '+r['pages']}, {r['year']}. https://doi.org/{r['doi']}")
p('[16] Computer Graphics Fundamentals course lecture collection supplied by the team, including Transformations, Viewing, Shading, Texture Mapping, Hierarchical Models and Procedural Methods. Slides credit E. Angel and D. Shreiner, Interactive Computer Graphics, 6th ed., Addison-Wesley, 2012. PDF-page mappings appear in Appendix A and docs/COURSE_MAP.md.')
p('[17] GLFW contributors, GLFW 3.4 documentation, https://www.glfw.org/docs/3.4/ . Accessed October 19, 2026.')
p('[18] Khronos Group, OpenGL 3.3 Core Profile Specification, https://registry.khronos.org/OpenGL/specs/gl/glspec33.core.pdf . API reference for buffers, shaders and drawing.')
p('[19] Matej Bel University Library, Scimago Journal & Country Rank: Computer Graphics and Computer-Aided Design, 2024 metrics (Scopus data as of March 2025), pp. 1 and 3. https://kniznica.umb.sk/app/cmsFile.php?ID=21899&disposition=i .')
p('Library credits and license locations: GLFW, zlib/libpng; GLAD, MIT and Khronos notices; GLM, MIT or Happy Bunny; Dear ImGui, MIT. Exact versions, authors, repository links and retained licenses are listed in CREDITS.md. No external models, texture photographs, sound tracks or copied shader snippets are used.')

new('9 AI Use Note')
p('OpenAI Codex provided substantial assistance in producing this project. It generated the initial C++ application, original procedural geometry and textures, GLSL shaders, simulation tests, build files, documentation, literature-review prose, report and presentation. It also assisted with compiling, debugging, searching and verifying references, measuring performance, and recording the demonstration.')
p('The user selected the Baikonur theme, supplied the assignment documents and course lectures, and provided the four team members\' names and groups. At preparation time, the team has not yet confirmed individual source changes, independent verification or actual module contributions. The package therefore does not attribute the generated implementation to particular students or claim a history of student commits.')
p('AI-assisted revisions made during validation included fixing the GLAD definition scope, correcting the emission ratio at high frame rates, moving the controls help to avoid overlap, separating fixed-resolution recording from window size, correcting framebuffer sample compatibility, and locking the scripted demo against unrelated input. These changes were followed by the automated and rendering checks reported in Section 6.')
p('Before submitting, the team must review the implementation, test it using the course setup, confirm or revise the explanations, and record its own substantive changes and responsibilities. Each member must be able to explain the whole pipeline and the mathematics relevant to their presentation. Any further AI-assisted modifications should be added to this note and the README with the specific work performed and how the team checked it.')
sub('Asset and source provenance')
p('Geometry, visual textures, shader effects and screenshots originate in the project. The report preserves the university logo from the supplied template. Scholarly papers informed comparisons and design reasoning; their advanced algorithms are not presented as implemented features. The source package retains library license files. The video contains real-time rendered application output with in-application captions and telemetry, encoded without speed alteration.')

new('Appendix A Course connections and defense')
table(['Course topic','Project connection'],[
('Transforms, L-4-1 pp. 9-10, 22-25','Column vectors, concatenation and local/parent transforms.'),
('Viewing, L-5-1 pp. 11-12, 16-22','LookAt, perspective projection and aspect ratio.'),
('Shading, L-5-2 pp. 29, 66-70','Modified Phong model and per-fragment evaluation.'),
('Textures, L-8 pp. 30-34, 38-40','UV interpolation, repeat sampling and mipmaps.'),
('Hierarchy, L-9 pp. 8-11, 29-32, 45','Parent-child composition and dynamic detachment.'),
('Procedural methods, L-10 pp. 14-22','Particle generation, attributes, motion and extinction.')],[2.8,3.7])
p('The references above use PDF page numbers. Course materials supplement the fifteen journal articles rather than replacing them. The detailed mapping also covers geometry, input, solid models and parametric surfaces.')
sub('Suggested presentation allocation')
p('Akbope Bakytkeldy (SE2426): slides 1-2, scope and requirements. Moldir Omar (SE2426): slides 3-4, pipeline and hierarchy. Ayaulym Talgat (SE2427): slides 5-7, shading and effects. Talant Altaikhan (SE2425): slides 8-10, interaction, tests and conclusions. This is a proposed speaking plan, not an assertion of completed development contributions.')
sub('Submission status')
p('The defense and submission dates remain to be confirmed. The team must verify topic registration, provide its hosted repository, add the TA, maintain genuine member commits, tag the actual submission revision v1.0, and upload the report PDF, slide deck and demo. The submission deadline is three days before defense at 23:59 Asia/Almaty. The local package cannot establish account-based submission or individual oral understanding.')

# Save authoring output, then preserve every original opaque ZIP part exactly.
tmp=root/'tmp/report-authored.docx';d.save(tmp)
final=root/'output/Baikonur_Report.docx'
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with zipfile.ZipFile(reference) as original,zipfile.ZipFile(tmp) as authored,zipfile.ZipFile(final,'w',zipfile.ZIP_DEFLATED) as out:
 for name in original.namelist():out.writestr(name,authored.read(name) if name in editable else original.read(name))
 for name in authored.namelist():
  if name not in original.namelist():out.writestr(name,authored.read(name))
with zipfile.ZipFile(reference) as original,zipfile.ZipFile(final) as output:
 assert all(original.read(n)==output.read(n) for n in original.namelist() if n not in editable)
 assert Document(reference).sections[0]._sectPr.xml==Document(final).sections[0]._sectPr.xml
print(final)
print('Expected pages: 19. Preserved original package parts outside editable slots.')
